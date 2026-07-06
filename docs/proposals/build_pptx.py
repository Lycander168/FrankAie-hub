#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""產生 2026-07-lycander-group-buy-deck.pptx（與同名 HTML 簡報內容一致）。

用法：python3 build_pptx.py
價格欄之〔提報前確認〕〔待填〕為預留文字，請於 PowerPoint 內直接編輯。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 視覺 tokens（同 HTML 版）─────────────────────────────
BG      = RGBColor(0x13, 0x1B, 0x2B)   # 夜空藍
CARD    = RGBColor(0x1A, 0x24, 0x38)
INK     = RGBColor(0xEE, 0xF1, 0xF7)
MUTED   = RGBColor(0x9A, 0xA5, 0xBC)
FAINT   = RGBColor(0x6B, 0x76, 0x90)
ACCENT  = RGBColor(0xE5, 0xB5, 0x4F)   # 月相金
ACCENT_D= RGBColor(0x8A, 0x65, 0x12)
HEADBG  = RGBColor(0x26, 0x2A, 0x2E)   # 表頭底
LINE    = RGBColor(0x3A, 0x44, 0x5C)
FONT    = "Microsoft JhengHei"

SW, SH = Inches(13.333), Inches(7.5)
PEND = "〔提報前確認〕"

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


def new_slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = BG
    r.line.fill.background(); r.shadow.inherit = False
    return s


def _set(para, text, size, color, bold=False, align=PP_ALIGN.LEFT, spacing=None):
    para.text = text
    para.alignment = align
    if spacing:
        para.space_after = Pt(spacing)
    for run in para.runs:
        f = run.font
        f.name = FONT; f.size = Pt(size); f.bold = bold
        f.color.rgb = color
        # 東亞字體（中文）需另設 a:ea，否則 PowerPoint 以預設中文字體顯示
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            latin = rPr.find(qn('a:latin'))
            ea = rPr.makeelement(qn('a:ea'), {})
            if latin is not None:
                latin.addnext(ea)
            else:
                rPr.append(ea)
        ea.set('typeface', FONT)


def textbox(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, color, bold[, align])"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        align = ln[4] if len(ln) > 4 else PP_ALIGN.LEFT
        _set(p, ln[0], ln[1], ln[2], ln[3], align, spacing=4)
    return tb


def footer(slide, n):
    textbox(slide, Inches(0.7), Inches(7.02), Inches(6), Inches(0.4),
            [("LYCANDER GROUP × 團購合作提案", 9, FAINT, False)])
    textbox(slide, Inches(12.2), Inches(7.02), Inches(0.9), Inches(0.4),
            [(f"{n:02d}", 9, FAINT, False, PP_ALIGN.RIGHT)])


def eyebrow(slide, text, y=Inches(0.55)):
    textbox(slide, Inches(0.7), y, Inches(8), Inches(0.4),
            [(text, 12, ACCENT, True)])


def title(slide, text, y=Inches(0.95), size=26):
    textbox(slide, Inches(0.7), y, Inches(12), Inches(1.0),
            [(text, size, INK, True)])


def section_tag(slide, text):
    tag = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(0.5),
                                 Inches(1.1), Inches(0.34))
    tag.fill.solid(); tag.fill.fore_color.rgb = ACCENT
    tag.line.fill.background(); tag.shadow.inherit = False
    tf = tag.text_frame; tf.margin_top = 0; tf.margin_bottom = 0
    _set(tf.paragraphs[0], text, 12, BG, True, PP_ALIGN.CENTER)


def card(slide, x, y, w, h, lines):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.adjustments[0] = 0.06
    c.fill.solid(); c.fill.fore_color.rgb = CARD
    c.line.color.rgb = LINE; c.line.width = Pt(0.75); c.shadow.inherit = False
    tf = c.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.16); tf.margin_bottom = Inches(0.14)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, ln[0], ln[1], ln[2], ln[3], spacing=4)
    return c


def kpi_row(slide, items, y=Inches(2.0), h=Inches(2.3)):
    x = Inches(0.7); gap = Inches(0.25)
    w = int((SW - Inches(1.4) - gap * (len(items) - 1)) / len(items))
    for i, (big, lbl) in enumerate(items):
        cx = x + i * (w + gap)
        card(slide, cx, y, w, h, [(big, 24, INK, True), (lbl, 11.5, MUTED, False)])
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, y, w, Pt(3))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background(); bar.shadow.inherit = False


def bullets(slide, x, y, w, h, items, size=13):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set(p, "•  " + t, size, MUTED, False, spacing=10)
    return tb


def table(slide, headers, rows, x, y, w, col_ratios=None, size=11, row_h=0.42):
    n_r, n_c = len(rows) + 1, len(headers)
    gt = slide.shapes.add_table(n_r, n_c, x, y, w, Inches(row_h * n_r)).table
    if col_ratios:
        total = sum(col_ratios)
        for i, r in enumerate(col_ratios):
            gt.columns[i].width = Emu(int(w * r / total))
    for ci, htxt in enumerate(headers):
        cell = gt.cell(0, ci)
        cell.fill.solid(); cell.fill.fore_color.rgb = HEADBG
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set(cell.text_frame.paragraphs[0], htxt, size, ACCENT, True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = gt.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = CARD
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            bold = (ci == 0)
            color = INK if bold else MUTED
            if PEND in str(val) or "待填" in str(val):
                color = ACCENT
            _set(cell.text_frame.paragraphs[0], str(val), size, color, bold)
    return gt


def moon(slide, x=Inches(10.2), y=Inches(0.8), d=Inches(2.1)):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    o.fill.gradient()
    stops = o.fill.gradient_stops
    stops[0].color.rgb = RGBColor(0xF3, 0xD6, 0x8C); stops[0].position = 0.0
    stops[1].color.rgb = ACCENT_D; stops[1].position = 1.0
    o.line.fill.background(); o.shadow.inherit = False


# ── 01 封面 ──────────────────────────────────────────────
s = new_slide(); moon(s)
textbox(s, Inches(0.9), Inches(1.15), Inches(8), Inches(0.5),
        [("L Y C A N D E R   G R O U P", 14, ACCENT, True)])
textbox(s, Inches(0.9), Inches(1.75), Inches(8.8), Inches(2.4),
        [("團購合作提案", 44, INK, True), ("四大主打商品 × 多入組方案", 44, INK, True)])
textbox(s, Inches(0.9), Inches(4.05), Inches(8.6), Inches(1.0),
        [("嘖嘖募資實績 × 國際品牌代理 × 授權親子商品——為全台團購公司與團購主準備的一站式開團提案。", 15, MUTED, False)])
for i, (k, v) in enumerate([("提案對象", "全台知名團購公司及團購主"), ("檔期", "2026 Q3–Q4"), ("聯絡窗口", "service@lycander.tw")]):
    textbox(s, Inches(0.9) + Inches(3.9) * i, Inches(5.5), Inches(3.8), Inches(0.9),
            [(k, 11, FAINT, False), (v, 13, INK, True)])
footer(s, 1)

# ── 02 提案摘要 ──────────────────────────────────────────
s = new_slide(); eyebrow(s, "EXECUTIVE SUMMARY"); title(s, "一頁看懂本次提案")
kpi_row(s, [
    ("4", "主打商品：小露娜／Wokyis M5／SHARGE Disk Pro／Nettec 兒童牙刷"),
    ("3＋1", "3C 科技 ＋ 親子生活，一次覆蓋兩大黃金客群"),
    ("嘖嘖", "三項商品皆有嘖嘖募資實績可直接引用開團"),
    ("多入組", "每品備妥單入＋多入組，另有跨產品主題組合"),
])
textbox(s, Inches(0.7), Inches(4.7), Inches(12), Inches(1.6),
        [("合作模式採「批售供貨」或「分潤抽成」雙軌，MOQ 與價格階梯詳見「合作條件」頁；"
          "所有標示〔提報前確認〕之欄位為保留給雙方議定的彈性空間。", 13, MUTED, False)])
footer(s, 2)

# ── 03 關於 LYCANDER ────────────────────────────────────
s = new_slide(); eyebrow(s, "ABOUT US"); title(s, "關於 LYCANDER")
bullets(s, Inches(0.7), Inches(2.0), Inches(6.2), Inches(4.2), [
    "自有品牌：LYCANDER（小露娜 LUNA Mag、OLIKA 系列充電產品）",
    "國際品牌台灣通路：Wokyis（M5 擴充座）、SHARGE 閃極（Disk Pro）",
    "授權合作：Nettec 優樂兒童電動牙刷",
    "既有通路：lycander.tw 官網、嘖嘖募資、PackUp；Nettec 已上架 momo／蝦皮／全家行動購／博客來／樂天",
])
card(s, Inches(7.2), Inches(2.0), Inches(5.4), Inches(3.4), [
    ("我們提供團購主", 15, INK, True),
    ("•  現貨供應與彈性出貨（代發／整批）", 12.5, MUTED, False),
    ("•  台灣在地保固與售後客服", 12.5, MUTED, False),
    ("•  完整行銷素材包與開團文案支援", 12.5, MUTED, False),
    ("•  開團期間專屬窗口即時對接", 12.5, MUTED, False),
])
footer(s, 3)

# ── 04 為什麼適合開團 ────────────────────────────────────
s = new_slide(); eyebrow(s, "WHY GROUP-BUY"); title(s, "為什麼這四支商品適合開團")
whys = [
    ("1｜募資實績可背書", "嘖嘖專案頁銷售數字可直接引用於開團文案——例：Wokyis M5 嘖嘖 807 人次贊助、2026/4/15 達標。"),
    ("2｜賣點一句話就能轉述", "「一秒展開的口袋充電站」「復古造型 13 合 1 擴充座」「全球首款主動散熱 PSSD」「迪士尼／三麗鷗授權牙刷」。"),
    ("3｜客群互補、雙線開團", "3C 打上班族與果粉，牙刷打媽媽社群；同一檔期可雙線操作，擴大觸及不互搶。"),
    ("4｜多入組拉高客單", "雙入／三入／家庭組湊單門檻低，符合團購「揪團湊件」習性，客單價與成團率同步提升。"),
]
for i, (h, b) in enumerate(whys):
    x = Inches(0.7) + (i % 2) * Inches(6.15)
    y = Inches(2.0) + (i // 2) * Inches(2.25)
    card(s, x, y, Inches(5.95), Inches(2.05), [(h, 14, INK, True), (b, 12, MUTED, False)])
footer(s, 4)

# ── 05 產品線總覽 ────────────────────────────────────────
s = new_slide(); eyebrow(s, "PRODUCT PORTFOLIO"); title(s, "產品線總覽")
table(s,
      ["產品", "定位", "目標客群", "建議售價（公開查證）", "銷售背書"],
      [
          ["LYCANDER 小露娜", "口袋型三合一無線充電站", "果粉、通勤、出差旅行", PEND, "嘖嘖 luna-mag"],
          ["Wokyis M5", "13合1 迷你螢幕擴充座", "Mac mini／桌面美學族", "標準版參考 NT$5,980", "嘖嘖 807 人次贊助"],
          ["SHARGE Disk Pro", "主動散熱 PSSD＋5合1 Hub", "攝影師、剪輯師", "美國官網 US$269 起", "嘖嘖＋國際媒體評測"],
          ["Nettec 兒童電動牙刷", "2–12 歲音波電動牙刷", "家長、親子社群", "NT$790／890／1,580", "momo／蝦皮／全家上架"],
      ],
      Inches(0.7), Inches(2.0), Inches(11.93), [1.1, 1.15, 1.0, 1.05, 1.0], size=11, row_h=0.62)
textbox(s, Inches(0.7), Inches(5.4), Inches(12), Inches(0.5),
        [("價格為公開通路查證之參考值；團購價與供貨價依「合作條件」頁價格階梯另議。", 10.5, FAINT, False)])
footer(s, 5)

# ── 06 小露娜 亮點 ──────────────────────────────────────
s = new_slide(); section_tag(s, "產品一")
title(s, "LYCANDER 小露娜（LUNA Mag）｜口袋型無線充電站——一秒展開，電就到", y=Inches(1.0), size=20)
bullets(s, Inches(0.7), Inches(2.2), Inches(6.4), Inches(3.8), [
    "三合一同充：手機＋耳機＋手錶一座搞定",
    "次世代 Qi2.2 無線充電標準，充電速度較前代提升約 70%",
    "口袋型摺疊設計，一秒展開即用；出差旅行免帶三條線",
    "LYCANDER 自有品牌，台灣在地保固與售後",
])
card(s, Inches(7.4), Inches(2.2), Inches(5.2), Inches(3.6), [
    ("完整規格", 15, INK, True),
    ("尺寸／重量／各埠瓦數配置：〔待填：以嘖嘖專案頁為準〕", 12.5, ACCENT, False),
    ("", 6, MUTED, False),
    ("嘖嘖專案：zeczec.com/projects/luna-mag", 11.5, MUTED, False),
    ("募資金額／贊助人數請於提報當日以專案頁即時數字截圖引用。", 11.5, FAINT, False),
])
footer(s, 6)

# ── 07 小露娜 價格 ──────────────────────────────────────
s = new_slide(); section_tag(s, "產品一")
title(s, "小露娜｜團購價格方案（多入組預留）", y=Inches(1.0))
table(s,
      ["方案", "內容", "建議售價", "團購價（建議 78–85 折）", "團購主毛利空間"],
      [
          ["單入", "LUNA Mag ×1", PEND, PEND, PEND],
          ["雙入組", "×2（自用＋送禮）", "單入 ×2", "雙入再降 3–5%", PEND],
          ["三入團購組", "×3（辦公室揪團）", "單入 ×3", "三入再降 5–8%", PEND],
      ],
      Inches(0.7), Inches(2.1), Inches(11.93), [0.8, 1.05, 0.95, 1.35, 1.0], size=12, row_h=0.6)
textbox(s, Inches(0.7), Inches(4.9), Inches(12), Inches(0.6),
        [("折扣區間為市場慣例之建議結構；實際團購底價由 LYCANDER 於提報前確認填入，開團期間享全通路價格保護。", 10.5, FAINT, False)])
footer(s, 7)

# ── 08 M5 亮點 ──────────────────────────────────────────
s = new_slide(); section_tag(s, "產品二")
title(s, "Wokyis M5｜13 合 1 迷你螢幕擴充座——復古外觀，裡面全是現代武裝", y=Inches(1.0), size=20)
bullets(s, Inches(0.7), Inches(2.2), Inches(6.4), Inches(3.4), [
    "復古經典電腦造型＋ 5 吋 HD 螢幕，桌面即是話題",
    "13 合 1 連接埠；專為 Mac mini 設計，相容 MacBook Air/Pro 與 Windows",
    "M.2 NVMe SSD 擴充槽（最高支援 8TB，內建散熱風扇）",
])
card(s, Inches(7.4), Inches(2.2), Inches(5.2), Inches(1.55), [
    ("標準版｜USB-C 10Gbps", 14, INK, True),
    ("日常高速外接與文書效率首選", 12, MUTED, False),
])
card(s, Inches(7.4), Inches(4.0), Inches(5.2), Inches(1.55), [
    ("專業版｜Thunderbolt 5 80Gbps", 14, INK, True),
    ("8K@60Hz DisplayPort 輸出，專業剪輯與高速 SSD 讀寫", 12, MUTED, False),
])
footer(s, 8)

# ── 09 M5 背書 ──────────────────────────────────────────
s = new_slide(); section_tag(s, "產品二")
title(s, "Wokyis M5｜銷售實績與背書", y=Inches(1.0))
kpi_row(s, [
    ("807", "嘖嘖贊助人次（796 位獨立贊助者）"),
    ("4/15", "2026 年嘖嘖達標日"),
    ("KS", "同步 Kickstarter 國際上線"),
    ("$5,980", "標準版市售參考價（NT，lycander.tw／PackUp 販售中）"),
], y=Inches(2.1), h=Inches(2.0))
textbox(s, Inches(0.7), Inches(4.5), Inches(12), Inches(1.2), [
    ("嘖嘖專案：zeczec.com/projects/wokyis（常見問答 /faqs）", 12, MUTED, False),
    ("總集資金額：〔待填：嘖嘖頁面即時數字〕——提報當日截圖引用，數字最有說服力。", 12, ACCENT, False),
])
footer(s, 9)

# ── 10 M5 價格 ──────────────────────────────────────────
s = new_slide(); section_tag(s, "產品二")
title(s, "Wokyis M5｜團購價格方案（多入組預留）", y=Inches(1.0))
table(s,
      ["方案", "內容", "建議售價", "團購價（建議 78–85 折）", "團購主毛利空間"],
      [
          ["標準版單入", "M5（10Gbps）×1", "NT$5,980（參考）", PEND, PEND],
          ["專業版單入", "M5（TB5 80Gbps）×1", PEND, PEND, PEND],
          ["雙入辦公組", "標準版 ×2", "NT$11,960（參考）", "雙入再降 3–5%", PEND],
          ["SSD 升級組", "M5 ＋ SSD（容量另議）", PEND, "組合價另議", PEND],
      ],
      Inches(0.7), Inches(2.1), Inches(11.93), [0.85, 1.15, 1.05, 1.35, 1.0], size=12, row_h=0.58)
footer(s, 10)

# ── 11 Disk Pro 亮點 ────────────────────────────────────
s = new_slide(); section_tag(s, "產品三")
title(s, "SHARGE Disk Pro｜全球首款主動散熱 PSSD＋5 合 1 Hub", y=Inches(1.0), size=20)
bullets(s, Inches(0.7), Inches(2.1), Inches(6.6), Inches(4.4), [
    "全球首款主動散熱行動固態硬碟：渦輪風扇 7,000–10,000 RPM 自動調速",
    "持續寫入 480GB 以上仍維持 50°C 以下，長時間傳輸不掉速",
    "10Gbps 高速傳輸；容量 1TB／2TB／4TB",
    "5 合 1 多埠 Hub（含 HDMI 2.1，支援 8K）",
    "150g 磁吸便攜；iPhone ProRes 直錄、iPad 即時備份、Mac/Windows 擴充皆適用",
], size=12.5)
card(s, Inches(7.6), Inches(2.1), Inches(5.0), Inches(3.4), [
    ("一機四用", 15, INK, True),
    ("儲存＋擴充＋外接螢幕＋供電——創作者出門只帶一顆。", 12.5, MUTED, False),
    ("", 6, MUTED, False),
    ("嘖嘖檔期曾祭出 VIP100 折扣碼＋HDMI 2.1 8K 認證線贈品，開團贈品操作有前例可循。", 11.5, FAINT, False),
])
footer(s, 11)

# ── 12 Disk Pro 背書 ────────────────────────────────────
s = new_slide(); section_tag(s, "產品三")
title(s, "SHARGE Disk Pro｜銷售實績與背書", y=Inches(1.0))
kpi_row(s, [
    ("首款", "全球第一款主動散熱 PSSD，品類開創者"),
    ("媒體", "Tom's Hardware（2TB 版）、PetaPixel 國際評測"),
    ("US$269", "美國官網起售價（1TB），國際定價可對照"),
    ("3 梯次", "嘖嘖出貨梯次 3/25、4/15、4/25，交付紀錄透明"),
], y=Inches(2.1), h=Inches(2.0))
textbox(s, Inches(0.7), Inches(4.5), Inches(12), Inches(1.2), [
    ("嘖嘖專案：zeczec.com/projects/disk-pro（贊助者留言 /comments 可作口碑引用）", 12, MUTED, False),
    ("募資金額／贊助人數：〔待填：嘖嘖頁面即時數字〕", 12, ACCENT, False),
])
footer(s, 12)

# ── 13 Disk Pro 價格 ────────────────────────────────────
s = new_slide(); section_tag(s, "產品三")
title(s, "SHARGE Disk Pro｜團購價格方案（多入組預留）", y=Inches(1.0))
table(s,
      ["方案", "內容", "建議售價", "團購價（建議 78–85 折）", "團購主毛利空間"],
      [
          ["1TB 單入", "Disk Pro 1TB ×1", PEND, PEND, PEND],
          ["2TB 單入", "Disk Pro 2TB ×1", PEND, PEND, PEND],
          ["4TB 單入", "Disk Pro 4TB ×1", PEND, PEND, PEND],
          ["創作者雙入組", "2TB ×2（工作＋備份）", PEND, "雙入再降 3–5%", PEND],
      ],
      Inches(0.7), Inches(2.1), Inches(11.93), [0.9, 1.15, 1.0, 1.35, 1.0], size=12, row_h=0.58)
textbox(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.5),
        [("台灣售價與各容量團購底價由 LYCANDER 依原廠供貨條件於提報前確認。", 10.5, FAINT, False)])
footer(s, 13)

# ── 14 Nettec 亮點 ──────────────────────────────────────
s = new_slide(); section_tag(s, "產品四")
title(s, "Nettec 授權兒童電動牙刷｜優樂 NETTEC——展開小朋友的牙齒冒險故事", y=Inches(1.0), size=20)
bullets(s, Inches(0.7), Inches(2.1), Inches(6.6), Inches(4.4), [
    "U 型矽膠刷頭音波電動牙刷，適用 2–12 歲，多段清潔模式",
    "迪士尼授權款：小熊維尼、三眼怪、史迪奇、獅子王",
    "三麗鷗授權款：酷洛米、大耳狗",
    "經典恐龍造型款（綠／粉）、輕巧攜帶型（黑／白／粉）",
    "已上架 momo、蝦皮、全家行動購、博客來、樂天——通路口碑可查",
], size=12.5)
card(s, Inches(7.6), Inches(2.1), Inches(5.0), Inches(3.2), [
    ("長線價值：耗材回購", 15, INK, True),
    ("首團賣主機、回購團賣替換刷頭——適合固定班表的團購社群，一支商品吃兩種檔期。", 12.5, MUTED, False),
    ("", 6, MUTED, False),
    ("官網：www.nettec.com.tw", 11.5, MUTED, False),
])
footer(s, 14)

# ── 15 Nettec 價格 ──────────────────────────────────────
s = new_slide(); section_tag(s, "產品四")
title(s, "Nettec 兒童牙刷｜團購價格方案（多入組預留）", y=Inches(1.0))
table(s,
      ["方案", "內容", "建議售價", "團購價（建議 78–85 折）", "團購主毛利空間"],
      [
          ["攜帶型單入", "攜帶型電動牙刷 ×1", "NT$790", PEND, PEND],
          ["恐龍款單入", "U 型恐龍造型 ×1", "NT$890", PEND, PEND],
          ["授權款單入", "迪士尼授權款 ×1", "NT$1,580", PEND, PEND],
          ["兄弟姊妹 2 入組", "恐龍款 ×2（雙色）", "NT$1,780", "2 入再降 3–5%", PEND],
          ["親子家庭組", "授權款 ×2 ＋ 替換刷頭組", PEND, "組合價另議", PEND],
          ["刷頭耗材加購", "替換刷頭（回購型）", "〔待填：官網耗材價〕", PEND, PEND],
      ],
      Inches(0.7), Inches(1.95), Inches(11.93), [1.0, 1.15, 1.05, 1.25, 0.95], size=11, row_h=0.52)
footer(s, 15)

# ── 16 跨產品多入組 ──────────────────────────────────────
s = new_slide(); eyebrow(s, "BUNDLE MATRIX"); title(s, "跨產品多入組總表（主題組合提報）")
table(s,
      ["主題組合", "內容", "訴求", "組合價"],
      [
          ["桌面工作站組", "Wokyis M5 ＋ Disk Pro", "Mac mini 桌面一次升級到位", PEND],
          ["出門帶電組", "小露娜 ×2", "家裡一座、包包一座", PEND],
          ["創作者全配組", "M5 專業版 ＋ Disk Pro 2TB ＋ 小露娜", "剪輯師／攝影師整套帶走", PEND],
          ["親子潔牙組", "Nettec 授權款 ×2 ＋ 刷頭耗材", "二寶家庭一次買齊", PEND],
          ["全家超值箱", "小露娜 ＋ Nettec ×2", "爸媽的 3C＋孩子的牙刷，一團全包", PEND],
      ],
      Inches(0.7), Inches(1.95), Inches(11.93), [0.85, 1.5, 1.35, 0.8], size=11, row_h=0.52)
textbox(s, Inches(0.7), Inches(5.3), Inches(12), Inches(0.6),
        [("組合折扣邏輯建議：組合價 ≦ 各品團購價加總再降 3–8%；以「送耗材／配件」取代破價，保護各品單售價格帶。", 10.5, FAINT, False)])
footer(s, 16)

# ── 17 合作條件 ──────────────────────────────────────────
s = new_slide(); eyebrow(s, "TERMS"); title(s, "合作條件（價格階梯）")
table(s,
      ["條件", "內容"],
      [
          ["合作模式", "A. 批售供貨（買斷）／ B. 分潤抽成（代收代付），依團購主慣例擇一"],
          ["價格階梯", f"建議售價 → 團購價（建議 78–85 折）→ 批發／供貨價 {PEND}"],
          ["MOQ", f"各品項最低開團量 {PEND}；多入組可合併計算"],
          ["價格保護（MAP）", "開團期間全通路不破團購價；檔期錯開既有通路促銷"],
          ["出貨", f"現貨品項下單後 X 個工作天出貨 {PEND}；可代發或整批出"],
          ["售後保固", "台灣在地保固；開團期間專屬客服窗口"],
          ["對帳", f"團後 X 天內對帳結算 {PEND}"],
      ],
      Inches(0.7), Inches(1.95), Inches(11.93), [0.6, 2.4], size=11, row_h=0.52)
footer(s, 17)

# ── 18 開團支援 ──────────────────────────────────────────
s = new_slide(); eyebrow(s, "SUPPORT"); title(s, "開團支援")
card(s, Inches(0.7), Inches(2.0), Inches(5.95), Inches(3.2), [
    ("素材支援", 15, INK, True),
    ("•  產品圖庫＋開箱短影音", 12.5, MUTED, False),
    ("•  開團文案包（含嘖嘖實績截圖授權）", 12.5, MUTED, False),
    ("•  FAQ 客服懶人包，降低小編負擔", 12.5, MUTED, False),
])
card(s, Inches(6.9), Inches(2.0), Inches(5.95), Inches(3.2), [
    ("活動支援", 15, INK, True),
    ("•  直播帶貨：出借樣品、品牌方連線說明", 12.5, MUTED, False),
    ("•  KOL 合作可談，分潤機制彈性", 12.5, MUTED, False),
    ("•  開團 SOP 對接：選品 → 定價 → 素材 → 上架 → 出貨 → 對帳", 12.5, MUTED, False),
])
footer(s, 18)

# ── 19 檔期建議 ──────────────────────────────────────────
s = new_slide(); eyebrow(s, "CALENDAR"); title(s, "2026 Q3–Q4 檔期建議")
table(s,
      ["檔期", "主推商品", "理由"],
      [
          ["8 月｜開學季", "Nettec 兒童牙刷、多入親子組", "開學前家長採買潮"],
          ["9 月｜新機潮", "小露娜、SHARGE Disk Pro", "iPhone 新機發表帶動配件與 ProRes 儲存需求"],
          ["11 月｜雙 11", "全品項＋跨產品組合", "年度最大檔，多入組主打"],
          ["12 月｜聖誕年末", "Wokyis M5、授權款牙刷", "造型商品送禮需求"],
      ],
      Inches(0.7), Inches(2.0), Inches(11.93), [0.8, 1.2, 1.6], size=12, row_h=0.6)
footer(s, 19)

# ── 20 資料來源 ──────────────────────────────────────────
s = new_slide(); eyebrow(s, "SOURCES"); title(s, "資料來源")
textbox(s, Inches(0.7), Inches(2.0), Inches(12), Inches(3.9), [
    ("嘖嘖｜LUNA Mag 小露娜：https://www.zeczec.com/projects/luna-mag", 12.5, MUTED, False),
    ("嘖嘖｜Wokyis M5：https://www.zeczec.com/projects/wokyis（常見問答 /faqs）", 12.5, MUTED, False),
    ("嘖嘖｜SHARGE Disk Pro：https://www.zeczec.com/projects/disk-pro（留言 /comments）", 12.5, MUTED, False),
    ("優樂 NETTEC 官網：https://www.nettec.com.tw", 12.5, MUTED, False),
    ("輔助來源：SHARGE 官網、Tom's Hardware 與 PetaPixel 評測、lycander.tw、PackUp、momo／蝦皮／全家行動購／博客來 商品頁", 12.5, MUTED, False),
    ("", 8, MUTED, False),
    ("引用原則：對外提報時，嘖嘖募資金額與贊助人數請以專案頁當日即時數字截圖為準；本簡報僅收錄已查證之公開資訊，未查證欄位一律標示待填。", 11, FAINT, False),
])
footer(s, 20)

# ── 21 聯絡頁 ────────────────────────────────────────────
s = new_slide(); moon(s)
textbox(s, Inches(0.9), Inches(1.9), Inches(8), Inches(0.5),
        [("L Y C A N D E R   G R O U P", 14, ACCENT, True)])
textbox(s, Inches(0.9), Inches(2.5), Inches(8.8), Inches(2.2),
        [("期待與您", 40, INK, True), ("在 2026 Q3 一起開團", 40, INK, True)])
for i, (k, v) in enumerate([("Email", "service@lycander.tw"), ("官網", "www.lycander.tw")]):
    textbox(s, Inches(0.9) + Inches(4.0) * i, Inches(5.3), Inches(3.8), Inches(0.9),
            [(k, 11, FAINT, False), (v, 14, INK, True)])
footer(s, 21)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "2026-07-lycander-group-buy-deck.pptx")
prs.save(out)
print(f"saved: {out} ({len(prs.slides._sldIdLst)} slides)")
